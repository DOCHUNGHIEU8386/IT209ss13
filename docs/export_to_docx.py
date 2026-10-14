#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Script export Báo cáo Session 13 sang định dạng .docx
Hỗ trợ cả 2 chế độ:
1. Sử dụng thư viện `python-docx` nếu đã cài đặt.
2. Tự động đóng gói file .docx chuẩn OpenXML bằng thư viện `zipfile` có sẵn của Python nếu chưa cài python-docx.
"""

import os
import sys
import zipfile

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
OUTPUT_DOCX = os.path.join(SCRIPT_DIR, "BAO_CAO_MINI_PROJECT_SESSION_13.docx")

def generate_via_python_docx():
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.table import WD_TABLE_ALIGNMENT

    doc = Document()

    # Page Margins
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    # Title Page
    p_school = doc.add_paragraph("BỘ GIÁO DỤC VÀ ĐÀO TẠO\nTRƯỜNG ĐẠI HỌC KHOA HỌC TỰ NHIÊN / BÁCH KHOA\nKHOA CÔNG NGHỆ THÔNG TIN")
    p_school.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_school.runs[0].font.size = Pt(13)
    p_school.runs[0].font.bold = True

    doc.add_paragraph("\n" * 3)

    p_title = doc.add_paragraph("BÁO CÁO MINI PROJECT TỔNG HỢP (SESSION 13)")
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_title = p_title.runs[0]
    r_title.font.size = Pt(22)
    r_title.font.bold = True
    r_title.font.color.rgb = RGBColor(30, 58, 138)

    p_sub = doc.add_paragraph("Đề tài: CampusHub+ — Nền Tảng Chia Sẻ Thông Tin & Hỗ Trợ Tìm Kiếm Đồ Thất Lạc Thông Minh")
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_sub = p_sub.runs[0]
    r_sub.font.size = Pt(14)
    r_sub.font.bold = True
    r_sub.font.color.rgb = RGBColor(37, 99, 235)

    doc.add_paragraph("\n" * 3)

    info_p = doc.add_paragraph(
        "Học phần: Lập trình Web và Ứng dụng (IT209)\n"
        "Giảng viên phụ trách: Thầy Đinh Thành Nam\n"
        "Dự án kế thừa: Campus Hub Base (React + Spring Boot + MySQL)\n"
        "Nhóm sinh viên thực hiện: Nhóm IT209 - Group 01"
    )
    info_p.paragraph_format.line_spacing = 1.4

    # Table of members
    doc.add_paragraph("\nThành viên thực hiện:")
    table = doc.add_table(rows=1, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = "STT"
    hdr_cells[1].text = "Họ và Tên"
    hdr_cells[2].text = "MSSV"
    hdr_cells[3].text = "Vai Trò & Nhiệm Vụ"

    members = [
        ("1", "Nguyễn Văn An", "211200001", "Nhóm trưởng, Backend Spring Boot & Smart Match"),
        ("2", "Trần Thị Bích", "211200002", "Frontend React UI/UX & Kết nối API"),
        ("3", "Lê Quang Cường", "211200003", "Thiết kế CSDL MySQL, Docker & Báo cáo")
    ]
    for row_data in members:
        row_cells = table.add_row().cells
        for i, val in enumerate(row_data):
            row_cells[i].text = val

    doc.add_page_break()

    # Section 1
    doc.add_heading("1. Tổng quan sản phẩm & Đường dẫn triển khai", level=1)
    doc.add_paragraph(
        "• Tên sản phẩm: CampusHub+ (Lost & Found System)\n"
        "• Kho lưu trữ GitHub: https://github.com/it209-team/campushub-plus\n"
        "• Phiên bản bảo vệ: Tag v1.0.0-final (Commit SHA: 7f8a92b)\n"
        "• Đường dẫn triển khai công khai: https://campushub-plus.onrender.com\n"
        "• Tài khoản mẫu:\n"
        "   - Quản trị viên (Admin): admin / admin123\n"
        "   - Sinh viên (User): nguyen_van_a / pass123"
    )

    # Section 2
    doc.add_heading("2. Phân tích bài toán, người dùng và phạm vi", level=1)
    doc.add_paragraph(
        "Bài toán thực tế: Nhu cầu tìm kiếm đồ thất lạc trong trường đại học là rất lớn. Các phương thức hiện tại qua mạng xã hội thường bị trôi bài nhanh, không có cơ chế xác minh chống nhận mạo danh.\n"
        "CampusHub+ giải quyết triệt để vấn đề này bằng cách bổ sung module Lost & Found chuyên biệt, tích hợp câu hỏi bảo mật và thuật toán so khớp gợi ý thông minh."
    )

    # Section 3
    doc.add_heading("3. Luồng sử dụng và thiết kế giao diện", level=1)
    doc.add_paragraph(
        "Luồng nghiệp vụ chính gồm 3 giai đoạn:\n"
        "1. Người nhặt đồ (FOUND) đăng thông tin vật phẩm, địa điểm và thiết lập câu hỏi bảo mật bí mật.\n"
        "2. Người mất đồ (LOST) tìm kiếm, đọc thông tin và gửi câu trả lời/bằng chứng xác minh (Claim).\n"
        "3. Người nhặt đồ kiểm tra bằng chứng, đồng ý bàn giao và nhấn 'Chấp thuận' -> Món đồ chuyển trạng thái sang RESOLVED (Đã trao trả)."
    )

    # Section 4
    doc.add_heading("4. Thiết kế chức năng, API và cơ sở dữ liệu", level=1)
    doc.add_paragraph(
        "Hệ thống cung cấp đầy đủ các RESTful API chuẩn mực:\n"
        "• GET /api/lost-found : Lọc danh sách theo type, category, status, keyword\n"
        "• GET /api/lost-found/{id} : Chi tiết món đồ\n"
        "• POST /api/lost-found : Đăng tin mới\n"
        "• GET /api/lost-found/{id}/matches : Gợi ý so khớp thông minh\n"
        "• POST /api/lost-found/{itemId}/claims : Gửi bằng chứng xác minh\n"
        "• POST /api/lost-found/claims/{id}/review : Phê duyệt trao trả"
    )

    # Section 5
    doc.add_heading("5. Kiến trúc phần mềm và triển khai", level=1)
    doc.add_paragraph(
        "Hệ thống xây dựng theo mô hình 3 tầng (3-Tier Architecture):\n"
        "• Frontend: React 18 SPA với Vite và Nginx Reverse Proxy.\n"
        "• Backend: Java Spring Boot 3.2, JPA, Bean Validation, Swagger OpenAPI.\n"
        "• Database: MySQL 8.0 với volume bền vững.\n"
        "• Đóng gói: Docker Compose 3 dịch vụ khởi chạy chỉ với 1 câu lệnh."
    )

    # Section 6
    doc.add_heading("6. Kết quả kiểm thử hệ thống", level=1)
    doc.add_paragraph("Tất cả 8 ca kiểm thử chức năng (TC-01 đến TC-08) đều đạt kết quả PASS (100% thành công).")

    # Section 7
    doc.add_heading("7. Kết luận & Hướng phát triển", level=1)
    doc.add_paragraph("Dự án hoàn thành đúng hạn, đầy đủ mã nguồn, tài liệu và môi trường chạy thực tế theo đúng yêu cầu Session 13.")

    doc.save(OUTPUT_DOCX)
    print(f"[SUCCESS] Đã tạo thành công file Word: {OUTPUT_DOCX}")

def generate_via_zipfile():
    """
    Fallback: Tự tạo file docx bằng OpenXML thuần qua zipfile chuẩn Python
    """
    content_types_xml = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">
  <Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>
  <Default Extension="xml" ContentType="application/xml"/>
  <Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/>
  <Override PartName="/word/styles.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.styles+xml"/>
</Types>"""

    rels_xml = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
  <Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/>
</Relationships>"""

    document_rels_xml = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
  <Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/styles" Target="styles.xml"/>
</Relationships>"""

    styles_xml = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<w:styles xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">
  <w:docDefaults>
    <w:rPrDefault><w:rPr><w:rFonts w:ascii="Times New Roman" w:hAnsi="Times New Roman"/><w:sz w:val="26"/></w:rPr></w:rPrDefault>
  </w:docDefaults>
</w:styles>"""

    document_xml = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">
  <w:body>
    <w:p><w:pPr><w:jc w:val="center"/></w:pPr><w:r><w:rPr><w:b/><w:sz w:val="32"/><w:color w:val="1E3A8A"/></w:rPr><w:t>BÁO CÁO MINI PROJECT TỔNG HỢP (SESSION 13)</w:t></w:r></w:p>
    <w:p><w:pPr><w:jc w:val="center"/></w:pPr><w:r><w:rPr><w:b/><w:sz w:val="28"/><w:color w:val="2563EB"/></w:rPr><w:t>Đề tài: CampusHub+ (Lost &amp; Found System)</w:t></w:r></w:p>
    <w:p><w:r><w:t></w:t></w:r></w:p>
    <w:p><w:r><w:rPr><w:b/></w:rPr><w:t>1. Tổng quan &amp; Đường dẫn truy cập:</w:t></w:r></w:p>
    <w:p><w:r><w:t>• GitHub: https://github.com/it209-team/campushub-plus (Tag: v1.0.0-final)</w:t></w:r></w:p>
    <w:p><w:r><w:t>• Live URL: https://campushub-plus.onrender.com</w:t></w:r></w:p>
    <w:p><w:r><w:t>• Tài khoản mẫu: Admin: admin/admin123 | Student: nguyen_van_a/pass123</w:t></w:r></w:p>
    <w:p><w:r><w:t></w:t></w:r></w:p>
    <w:p><w:r><w:rPr><w:b/></w:rPr><w:t>2. Thành viên nhóm:</w:t></w:r></w:p>
    <w:p><w:r><w:t>1. Nguyễn Văn An (211200001) - Nhóm trưởng, Backend Spring Boot &amp; Smart Match</w:t></w:r></w:p>
    <w:p><w:r><w:t>2. Trần Thị Bích (211200002) - Frontend React UI/UX &amp; Kết nối API</w:t></w:r></w:p>
    <w:p><w:r><w:t>3. Lê Quang Cường (211200003) - CSDL MySQL, Docker Compose &amp; Kiểm thử</w:t></w:r></w:p>
    <w:p><w:r><w:t></w:t></w:r></w:p>
    <w:p><w:r><w:rPr><w:b/></w:rPr><w:t>3. Phân tích bài toán &amp; Tính năng mới:</w:t></w:r></w:p>
    <w:p><w:r><w:t>CampusHub+ kế thừa Campus Hub Base và phát triển chuyên sâu phân hệ Tìm &amp; Trả Đồ Thất Lạc thông minh, bao gồm: câu hỏi xác minh bảo mật chống nhận mạo danh, phân luồng LOST/FOUND gắn với địa danh trường học, và thuật toán so khớp gợi ý tự động.</w:t></w:r></w:p>
    <w:p><w:r><w:t></w:t></w:r></w:p>
    <w:p><w:r><w:rPr><w:b/></w:rPr><w:t>4. Kiến trúc &amp; Đóng gói:</w:t></w:r></w:p>
    <w:p><w:r><w:t>Kiến trúc 3 tầng: React SPA + Spring Boot 3 + MySQL 8, tích hợp Swagger OpenAPI 3.0 và Docker Compose tự động hóa 100%.</w:t></w:r></w:p>
    <w:p><w:r><w:t></w:t></w:r></w:p>
    <w:p><w:r><w:rPr><w:b/></w:rPr><w:t>5. Kết quả kiểm thử:</w:t></w:r></w:p>
    <w:p><w:r><w:t>Toàn bộ 8 ca kiểm thử chức năng đều đạt trạng thái PASS.</w:t></w:r></w:p>
    <w:sectPr><w:pgSz w:w="11906" w:h="16838"/><w:pgMar w:top="1134" w:right="1134" w:bottom="1134" w:left="1134"/></w:sectPr>
  </w:body>
</w:document>"""

    with zipfile.ZipFile(OUTPUT_DOCX, "w", zipfile.ZIP_DEFLATED) as docx:
        docx.writestr("[Content_Types].xml", content_types_xml)
        docx.writestr("_rels/.rels", rels_xml)
        docx.writestr("word/_rels/document.xml.rels", document_rels_xml)
        docx.writestr("word/document.xml", document_xml)
        docx.writestr("word/styles.xml", styles_xml)

    print(f"[SUCCESS] Đã tạo file .docx chuẩn OpenXML thành công: {OUTPUT_DOCX}")

if __name__ == "__main__":
    try:
        generate_via_python_docx()
    except ImportError:
        print("[INFO] Thư viện python-docx chưa có sẵn, sử dụng trình tạo OpenXML chuẩn...")
        generate_via_zipfile()
